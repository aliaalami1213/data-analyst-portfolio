"""Step 3 of the model: recalculate the workbook, run integrity checks, export results.

openpyxl writes formulas but cannot calculate them, so this script recalculates
sony_financial_model.xlsx with the `formulas` library (an independent Excel-formula
engine), checks that the three statements tie together, and saves the headline
numbers and charts used in the README.

Writes  results/scenario_summary.csv, results/sensitivity.csv, results/base_case_forecast.csv
        images/02_valuation_by_scenario.png, images/03_sensitivity_heatmap.png,
        images/04_base_case_forecast.png

Usage:  python verify_model.py
"""
import re
import warnings
from pathlib import Path

import formulas
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

warnings.filterwarnings("ignore")
HERE = Path(__file__).parent
BOOK = HERE / "sony_financial_model.xlsx"
(HERE / "results").mkdir(exist_ok=True)

# --- recalculate every formula ------------------------------------------------------
solution = formulas.ExcelModel().loads(str(BOOK)).finish().calculate()
values = {}
for ref, cell in solution.items():
    match = re.match(r"'\[.*\](.+)'!([A-Z]+\d+)$", ref)
    if match:
        value = cell.value
        values[(match.group(1).upper(), match.group(2))] = value[0, 0] if hasattr(value, "shape") else value

wb = load_workbook(BOOK)


def val(sheet, ref):
    """Calculated value of a formula cell, or the stored value of an input cell."""
    return values.get((sheet.upper(), ref), wb[sheet][ref].value)


def row_of(sheet, label, start=1):
    for row in range(start, wb[sheet].max_row + 1):
        if wb[sheet][f"A{row}"].value == label:
            return row
    raise KeyError(f"{label!r} not found on {sheet}")


# --- integrity checks ---------------------------------------------------------------
n_formulas = sum(isinstance(c.value, str) and c.value.startswith("=")
                 for ws in wb.worksheets for row in ws.iter_rows() for c in row)
errors = [k for k, v in values.items() if isinstance(v, str) and v.startswith("#")]
assert not errors, f"Formula errors: {errors[:10]}"
print(f"{n_formulas:,} formulas recalculated, 0 errors")

FCST = "FGHIJ"
for name in ("Base", "Upside", "Downside"):
    check, cash_bs = row_of(name, "Balance check (assets less liabilities and equity; must be zero)"), \
        row_of(name, "Cash and cash equivalents")
    cash_cf = row_of(name, "Cash at end of year")
    ni, ni_cf = row_of(name, "Net income from continuing operations"), \
        row_of(name, "Net income from continuing operations", start=cash_bs)
    for col in FCST:
        assert abs(val(name, f"{col}{check}")) < 1, f"{name} {col}: balance sheet does not balance"
        assert abs(val(name, f"{col}{cash_bs}") - val(name, f"{col}{cash_cf}")) < 1, f"{name} {col}: cash does not tie"
        assert abs(val(name, f"{col}{ni}") - val(name, f"{col}{ni_cf}")) < 1, f"{name} {col}: net income does not tie"
    # reported operating income must be reproduced from its components
    ebit = row_of(name, "Operating income")
    reported = [1035255, 1276635, 1447507]
    assert [round(val(name, f"{c}{ebit}")) for c in "CDE"] == reported, f"{name}: historical operating income mismatch"
print("Checks passed: balance sheet balances, cash ties to cash flow, net income ties, "
      "reported operating income reproduced (all three scenarios, all five forecast years)")

# --- export results -----------------------------------------------------------------
first = row_of("Scenarios", "Growth adjustment (points per year)")
last = row_of("Scenarios", "Balance check, all forecast years (must be zero)")
summary = pd.DataFrame(
    {name: [val("Scenarios", f"{col}{r}") for r in range(first, last + 1)]
     for name, col in zip(("Base", "Upside", "Downside"), "BCD")},
    index=[wb["Scenarios"][f"A{r}"].value for r in range(first, last + 1)],
)
summary.index.name = "metric"
summary.to_csv(HERE / "results" / "scenario_summary.csv", float_format="%.4f")

head = row_of("Scenarios", "WACC  \\  terminal growth")
growth = [val("Scenarios", f"{c}{head}") for c in "BCDEF"]
waccs = [val("Scenarios", f"A{r}") for r in range(head + 1, head + 6)]
grid = pd.DataFrame([[val("Scenarios", f"{c}{r}") for c in "BCDEF"] for r in range(head + 1, head + 6)],
                    index=[f"{w:.1%}" for w in waccs], columns=[f"{g:.1%}" for g in growth])
grid.index.name = "WACC \\ terminal growth"
grid.round(0).to_csv(HERE / "results" / "sensitivity.csv")

lines = ["Sales", "Operating income", "Net income from continuing operations", "EBITDA",
         "Cash from operating activities", "Capital expenditure", "Free cash flow to the firm",
         "Cash and cash equivalents", "Total assets", "Total equity"]
years = [wb["Base"][f"{c}4"].value for c in "CDEFGHIJ"]
base = pd.DataFrame({y: [val("Base", f"{c}{row_of('Base', l)}") for l in lines] for y, c in zip(years, "CDEFGHIJ")},
                    index=lines)
base.index.name = "¥ millions"
base.to_csv(HERE / "results" / "base_case_forecast.csv", float_format="%.0f")

print("\nScenario summary:")
print(summary.to_string(float_format=lambda x: f"{x:,.3f}" if abs(x) < 10 else f"{x:,.0f}"))
print("\nValue per share (yen): WACC rows x terminal growth columns")
print(grid.round(0).astype(int).to_string())

# --- charts -------------------------------------------------------------------------
per_share = summary.loc["Value per share (¥)"]
market = summary.loc["Market price per share (¥)", "Base"]
fig, ax = plt.subplots(figsize=(7.5, 4.5))
bars = ax.bar(["Downside", "Base", "Upside"], [per_share["Downside"], per_share["Base"], per_share["Upside"]],
              color=["#C0504D", "#1F3864", "#4F9D69"], width=0.55)
ax.axhline(market, color="black", linestyle="--", linewidth=1)
ax.text(2.42, market, f"Market price\n¥{market:,.0f}", va="center", ha="left", fontsize=9)
for bar in bars:
    h = bar.get_height()
    ax.text(bar.get_x() + bar.get_width() / 2, h + 60, f"¥{h:,.0f}\n({h / market - 1:+.0%})", ha="center", fontsize=10,
            bbox=dict(facecolor="white", edgecolor="none", pad=1.5))
ax.set_ylim(0, per_share.max() * 1.22)
ax.set_xlim(-0.5, 3.1)
ax.set_ylabel("Value per share (¥)")
ax.set_title("Sony DCF value per share by scenario vs. market price")
for side in ("top", "right"):
    ax.spines[side].set_visible(False)
fig.tight_layout()
fig.savefig(HERE / "images" / "02_valuation_by_scenario.png", dpi=150)

fig, ax = plt.subplots(figsize=(7.5, 4.5))
ratio = grid.to_numpy(dtype=float) / market - 1
im = ax.imshow(ratio, cmap="RdYlGn", vmin=-0.6, vmax=0.6, aspect="auto")
ax.set_xticks(range(5))
ax.set_xticklabels(grid.columns)
ax.set_yticks(range(5))
ax.set_yticklabels(grid.index)
ax.set_xlabel("Terminal growth rate")
ax.set_ylabel("WACC")
for i in range(5):
    for j in range(5):
        ax.text(j, i, f"¥{grid.iloc[i, j]:,.0f}\n{ratio[i, j]:+.0%}", ha="center", va="center", fontsize=9,
                fontweight="bold" if (i, j) == (2, 2) else "normal")
ax.set_title(f"Base-case value per share and gap to market price (¥{market:,.0f})")
fig.colorbar(im, ax=ax, format=lambda x, _: f"{x:+.0%}", label="vs. market price")
fig.tight_layout()
fig.savefig(HERE / "images" / "03_sensitivity_heatmap.png", dpi=150)

fig, ax = plt.subplots(figsize=(9, 4.8))
x = np.arange(len(years))
tn = base / 1e6
colors = ["#8FAADC" if y.endswith("A") else "#1F3864" for y in years]
ax.bar(x, tn.loc["Sales"], color=colors, width=0.6, label="Sales (left)")
ax.set_xticks(x)
ax.set_xticklabels(years)
ax.set_ylabel("Sales (¥ trillion)")
ax2 = ax.twinx()
ax2.plot(x, tn.loc["Operating income"], color="#C55A11", marker="o", label="Operating income (right)")
ax2.plot(x[3:], tn.loc["Free cash flow to the firm"].iloc[3:], color="#4F9D69", marker="s",
         label="Free cash flow (right)")
ax2.set_ylabel("Operating income and free cash flow (¥ trillion)")
ax2.set_ylim(0, 2.2)
handles = ax.get_legend_handles_labels()[0] + ax2.get_legend_handles_labels()[0]
ax.legend(handles, [h.get_label() for h in handles], loc="upper left", frameon=False, fontsize=9)
ax.set_title("Base case: reported FY2023–FY2025 (light) and forecast FY2026–FY2030 (dark)")
ax.spines["top"].set_visible(False)
ax2.spines["top"].set_visible(False)
fig.tight_layout()
fig.savefig(HERE / "images" / "04_base_case_forecast.png", dpi=150)
print("\nSaved results/ and images/")
